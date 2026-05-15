"""
CLI 入口：批量坐标格式转换。

用法:
    # 单点转换
    python -m GIS_Tools.coord_transform.cli --from wgs84 --to gcj02 -lng 116.397 --lat 39.908

    # CSV 批量转换
    python -m GIS_Tools.coord_transform.cli --from gcj02 --to wgs84 -i input.csv -o output.csv --lng-col 经度 --lat-col 纬度

    # Excel 批量转换
    python -m GIS_Tools.coord_transform.cli --from wgs84 --to bd09 -i input.xlsx -o output.xlsx --lng-col lng --lat-col lat
"""

import argparse
import csv
import sys
from pathlib import Path

from .transform import TRANSFORMS


def _detect_format(path: str) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in (".xlsx", ".xls"):
        return "excel"
    return "csv"


def _read_csv(path: str) -> list[dict]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def _write_csv(path: str, rows: list[dict], fieldnames: list[str]):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _read_excel(path: str) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:]]


def _write_excel(path: str, rows: list[dict], fieldnames: list[str]):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(fieldnames)
    for row in rows:
        ws.append([row.get(f, "") for f in fieldnames])
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="坐标转换工具 — WGS84/GCJ02/BD09 互转")
    parser.add_argument("--from", dest="from_crs", choices=["wgs84", "gcj02", "bd09"], required=True)
    parser.add_argument("--to", choices=["wgs84", "gcj02", "bd09"], required=True)
    parser.add_argument("-i", "--input", help="输入文件路径 (.csv / .xlsx)")
    parser.add_argument("-o", "--output", help="输出文件路径")
    parser.add_argument("--lng-col", default="lng", help="经度列名，默认 lng")
    parser.add_argument("--lat-col", default="lat", help="纬度列名，默认 lat")
    parser.add_argument("-lng", "--lng", type=float, help="单点经度")
    parser.add_argument("--lat", type=float, help="单点纬度")
    args = parser.parse_args()

    from_crs = args.from_crs
    to_crs = args.to
    key = f"{from_crs}_to_{to_crs}"

    transform = TRANSFORMS.get(key)
    if transform is None:
        print(f"不支持的转换方向: {from_crs} → {to_crs}", file=sys.stderr)
        print(f"支持的转换: {', '.join(TRANSFORMS)}", file=sys.stderr)
        sys.exit(1)

    # 单点模式
    if args.lng is not None and args.lat is not None:
        new_lng, new_lat = transform(args.lng, args.lat)
        print(f"{from_crs}: ({args.lng}, {args.lat})")
        print(f"{to_crs}: ({new_lng:.6f}, {new_lat:.6f})")
        return

    # 文件模式
    if not args.input or not args.output:
        parser.error("文件模式需要 -i 和 -o，单点模式需要 --lng 和 --lat")

    fmt = _detect_format(args.input)
    read_fn = _read_excel if fmt == "excel" else _read_csv
    write_fn = _write_excel if _detect_format(args.output) == "excel" else _write_csv

    rows = read_fn(args.input)
    lng_col = args.lng_col
    lat_col = args.lat_col

    if lng_col not in rows[0] or lat_col not in rows[0]:
        print(f"列名 '{lng_col}' 或 '{lat_col}' 不存在。文件列名: {list(rows[0])}", file=sys.stderr)
        sys.exit(1)

    fieldnames = list(rows[0].keys())
    for row in rows:
        try:
            new_lng, new_lat = transform(float(row[lng_col]), float(row[lat_col]))
            row[lng_col] = f"{new_lng:.6f}"
            row[lat_col] = f"{new_lat:.6f}"
        except (ValueError, TypeError):
            print(f"跳过无效坐标行: {row}", file=sys.stderr)

    write_fn(args.output, rows, fieldnames)
    print(f"转换完成: {len(rows)} 行 → {args.output}")


if __name__ == "__main__":
    main()
